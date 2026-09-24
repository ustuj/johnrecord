from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

root = Path(__file__).resolve().parents[1]
screenshots = root / 'docs' / 'ui_check'
out = root / 'docs' / 'module2_screenshots.docx'

items = [
    ('01_login.png', 'Окно входа'),
    ('02_login_error.png', 'Ошибка авторизации с информативным сообщением'),
    ('03_guest.png', 'Каталог гостя'),
    ('04_client.png', 'Каталог авторизированного клиента'),
    ('05_manager_controls.png', 'Каталог менеджера: поиск, поставщик и сортировка'),
    ('06_manager_orders.png', 'Заказы менеджера'),
    ('07_admin_cards.png', 'Каталог администратора: товары, скидки и остатки'),
    ('08_product_new.png', 'Форма добавления пластинки'),
    ('09_product_edit.png', 'Форма редактирования пластинки с ID только для чтения'),
    ('10_order_new.png', 'Форма добавления заказа'),
    ('11_forbidden_delete.png', 'Предупреждение при попытке удалить пластинку из заказа'),
    ('12_cart.png', 'Корзина с выбранной пластинкой'),
]


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.top_margin = Cm(0.45)
section.bottom_margin = Cm(0.45)
section.left_margin = Cm(0.65)
section.right_margin = Cm(0.65)

styles = doc.styles
styles['Normal'].font.name = 'Times New Roman'
styles['Normal'].font.size = Pt(10)

# Cover page.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(45)
r = p.add_run('USTUJ RECORDS — Доказательства отладки и корректной работы')
r.bold = True
r.font.name = 'Times New Roman'
r.font.size = Pt(22)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
r = p.add_run('Модуль 2–4 · демонстрационный проект по КОД 09.02.07-2-2026')
r.font.name = 'Times New Roman'
r.font.size = Pt(13)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(25)
r = p.add_run('Скриншоты собраны после автоматической проверки основных сценариев авторизации, ролей, каталога, CRUD товаров и заказов.')
r.font.name = 'Times New Roman'
r.font.size = Pt(11)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Дата проверки: 25.09.2026')
r.font.name = 'Times New Roman'
r.font.size = Pt(10)

doc.add_page_break()

for index, (filename, caption) in enumerate(items, 1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f'Модуль · {caption}')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)

    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(screenshots / filename), width=Cm(25.0))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f'Скриншот {index}. {caption}.')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(8)

    if index != len(items):
        doc.add_page_break()

# Footer.
for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    r = footer.add_run('USTUJ RECORDS · демонстрационный проект')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(7)

out.parent.mkdir(exist_ok=True)
doc.save(out)
print(out)
